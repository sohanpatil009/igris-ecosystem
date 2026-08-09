from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
ACCENT = RGBColor(0x3B, 0x82, 0xF6)
DARK = RGBColor(0x11, 0x18, 0x27)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
SECONDARY = RGBColor(0x64, 0x74, 0x8B)
SUCCESS = RGBColor(0x10, 0xB9, 0x81)


def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, font_size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = align
    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_title_bar(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.2), DARK)
    add_rect(slide, Inches(0), Inches(1.2), prs.slide_width, Inches(0.06), ACCENT)
    add_text(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6), title, font_size=28, color=LIGHT, bold=True)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.4), subtitle, font_size=14, color=SECONDARY)


def add_footer(slide, page_num, total=9):
    add_text(slide, Inches(0.5), Inches(7.1), Inches(12), Inches(0.3),
             f"IGRIS - Cross-Platform Device Ecosystem    |    Page {page_num}/{total}",
             font_size=10, color=SECONDARY)
