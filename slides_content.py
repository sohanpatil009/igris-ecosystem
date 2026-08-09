# slides_content.py - Slides for IGRIS presentation
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, font_size=18, color=None, bold=False, align=PP_ALIGN.LEFT):
    if color is None:
        color = RGBColor(0x11, 0x18, 0x27)
    text = (text.replace('[bull]', '\u25CF')
                .replace('[rarr]', '\u2192')
                .replace('[check]', '\u2714')
                .replace('[arrow]', '\u27A4')
                .replace('[diams]', '\u25C6'))
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = align
    return txBox

def add_rounded(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_title_bar(slide, title, subtitle=None):
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide_width = Inches(13.333)
    add_rect(slide, Inches(0), Inches(0), slide_width, Inches(1.2), DARK)
    add_rect(slide, Inches(0), Inches(1.2), slide_width, Inches(0.06), ACCENT)
    add_text(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6), title, font_size=28, color=LIGHT, bold=True)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.4), subtitle, font_size=14, color=SECONDARY)

def add_footer(slide, page_num, total=9):
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    add_text(slide, Inches(0.5), Inches(7.1), Inches(12), Inches(0.3),
             f"IGRIS - Cross-Platform Device Ecosystem    |    Page {page_num}/{total}",
             font_size=10, color=SECONDARY)


def slide_title(prs):
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK)
    add_rect(slide, Inches(0), Inches(2.5), prs.slide_width, Inches(0.08), ACCENT)
    add_text(slide, Inches(0.5), Inches(1.0), Inches(12), Inches(0.8),
             "IGRIS", font_size=60, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.5), Inches(1.8), Inches(12), Inches(0.6),
             "Cross-Platform Device Ecosystem", font_size=32, color=LIGHT, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(4.5), Inches(3.5), Inches(4.33), Inches(0.05), ACCENT)
    items = ["Universal Clipboard Sync", "Cross-Device Notification Mirror",
             "FastSwap File Transfer (P2P, TLS-encrypted)", "LAN Device Mesh with TLS Encryption"]
    for i, item in enumerate(items):
        add_text(slide, Inches(3.5), Inches(3.8 + i * 0.45), Inches(6.33), Inches(0.4),
                 f"[bull]  {item}", font_size=16, color=RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.5), Inches(6.5), Inches(12), Inches(0.4),
             "A Rust-powered, serverless, account-free device mesh for your LAN",
             font_size=14, color=SECONDARY, align=PP_ALIGN.CENTER)
    return slide


def slide_problem(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Problem Statement", "Why IGRIS Exists")
    problems = [
        "Fragmented device ecosystems force users to rely on cloud servers for basic cross-device tasks",
        "Existing solutions (AirDrop, Nearby Share) are vendor-locked and platform-specific",
        "Clipboard sync tools lack security - content transmitted in plaintext or via third-party servers",
        "Notification management is siloed - replies must originate from the source device",
        "File transfers between heterogeneous devices (Windows + macOS + Linux) remain cumbersome",
        "Privacy concerns: most sync tools require accounts, internet, and store data on remote servers"
    ]
    for i, prob in enumerate(problems):
        row = i // 2
        col = i % 2
        left = Inches(0.5 + col * 6.15)
        top = Inches(1.7 + row * 1.1)
        add_rounded(slide, left, top, Inches(5.9), Inches(0.95),
                    RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE2, 0xE8, 0xF0))
        add_text(slide, left + Inches(0.2), top + Inches(0.1), Inches(5.4), Inches(0.8),
                 f"{i+1}.  {prob}", font_size=12, color=DARK)
    add_footer(slide, 2)
    return slide


def slide_scope(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Scope & Objectives", "What IGRIS Aims to Achieve")
    add_text(slide, Inches(0.5), Inches(1.6), Inches(6), Inches(0.4), "PROJECT SCOPE", font_size=16, color=ACCENT, bold=True)
    scope_items = [
        "Build a cross-platform desktop app (Windows, macOS, Linux) as a single native binary",
        "Operate entirely on LAN - no cloud, no account, no internet required",
        "Secure peer-to-peer communication over TLS 1.3 with self-signed certificates",
        "Support device discovery, pairing, clipboard sync, notification sync, and file transfer",
        "Provide an optional voice assistant (offline STT/TTS + hybrid online mode)"
    ]
    for i, item in enumerate(scope_items):
        add_text(slide, Inches(0.7), Inches(2.1 + i * 0.42), Inches(5.8), Inches(0.4),
                 f"[rarr]  {item}", font_size=12, color=DARK)
    add_text(slide, Inches(7), Inches(1.6), Inches(6), Inches(0.4), "KEY OBJECTIVES", font_size=16, color=ACCENT, bold=True)
    obj_items = [
        "Zero-config device discovery on local subnet via IPv4 scan",
        "Loop-proof clipboard sync using SHA-256 hashing and two-marker diff",
        "Trust-gated data flow - only linked devices receive clipboard/notifications",
        "FastSwap file transfer with approval flow and live progress",
        "Extensible plugin architecture for voice interface and future features"
    ]
    for i, item in enumerate(obj_items):
        add_text(slide, Inches(7.2), Inches(2.1 + i * 0.42), Inches(5.8), Inches(0.4),
                 f"[rarr]  {item}", font_size=12, color=DARK)
    add_footer(slide, 3)
    return slide


def slide_literature(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Literature Review", "Existing Solutions & Comparison")
    add_text(slide, Inches(0.5), Inches(1.6), Inches(12.33), Inches(0.35),
             "Existing Cross-Device Ecosystem Solutions", font_size=14, color=DARK, bold=True)
    solutions = [
        ("Solution", "Platform", "Encryption", "Account", "Internet", "Open Source"),
        ("Apple AirDrop", "Apple only", "TLS", "Apple ID", "Not required", "No"),
        ("KDE Connect", "Win/Mac/Linux", "TLS", "None", "Not required", "Yes"),
        ("LocalSend", "All platforms", "TLS", "None", "Not required", "Yes"),
        ("Syncthing", "All platforms", "TLS", "None", "Optional", "Yes"),
        ("Pushbullet", "All platforms", "TLS", "Required", "Required", "No"),
        ("IGRIS (Ours)", "Win/Mac/Linux", "TLS 1.3", "None", "Not required", "MIT License")
    ]
    for row_idx, row_data in enumerate(solutions):
        top = Inches(2.0 + row_idx * 0.55)
        is_header = row_idx == 0
        is_last = row_idx == len(solutions) - 1
        if is_header:
            bg_color = DARK
        elif is_last:
            bg_color = RGBColor(0xDB, 0xEA, 0xFE)
        elif row_idx % 2 == 0:
            bg_color = RGBColor(0xF1, 0xF5, 0xF9)
        else:
            bg_color = RGBColor(0xFF, 0xFF, 0xFF)
        left = Inches(0.5)
        widths = [Inches(2.2), Inches(2.2), Inches(1.8), Inches(1.5), Inches(1.5), Inches(2.2)]
        for col_idx, (cell_text, width) in enumerate(zip(row_data, widths)):
            add_rect(slide, left, top, width, Inches(0.5), bg_color)
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.5))
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            cell.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            cell.line.width = Pt(0.5)
            add_text(slide, left + Inches(0.1), top + Inches(0.08), width - Inches(0.2), Inches(0.35),
                     cell_text, font_size=10, color=LIGHT if is_header else DARK, bold=is_header)
            left += width
    add_text(slide, Inches(0.5), Inches(6.0), Inches(12.33), Inches(0.35),
             "Key Finding: No existing solution combines a zero-account LAN mesh + optional voice interface + cross-platform native binary.",
             font_size=11, color=PRIMARY, bold=True)
    add_text(slide, Inches(0.5), Inches(6.4), Inches(12.33), Inches(0.35),
             "IGRIS fills this gap with a single Rust codebase that runs on every major desktop OS.",
             font_size=11, color=SECONDARY)
    add_footer(slide, 4)
    return slide


def slide_architecture(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Proposed System Architecture", "Layered Design of IGRIS")
    add_text(slide, Inches(0.5), Inches(1.6), Inches(6), Inches(0.4), "SYSTEM ARCHITECTURE (6 Layers)", font_size=16, color=ACCENT, bold=True)
    arch_layers = [
        ("Layer 1: Desktop UI (Dioxus 0.7)", "Panels: Devices, Notifications, FastSwap, Camera, File Search"),
        ("Layer 2: Device Mesh Runtime (src/eco/)", "Discovery, Clipboard, Notifications, Sync, Transport, Pairing"),
        ("Layer 3: FastSwap Module", "FastSwap: Server, Client, TLS, Approval Flow, Progress"),
        ("Layer 4: Voice Assistant & Plugins", "Optional voice: SenseVoice STT, Piper TTS, NLU, command processor"),
        ("Layer 5: Platform Abstraction", "Win32 FFI, macOS NotificationCenter, Linux D-Bus, xclip"),
        ("Layer 6: Networking Layer", "Axum HTTP Server (53327), TLS Proxy (53328), REST API"),
    ]
    for i, (layer, desc) in enumerate(arch_layers):
        top = Inches(2.1 + i * 0.65)
        fill = RGBColor(0xDB, 0xEA, 0xFE) if i % 2 == 0 else RGBColor(0xBD, 0xE0, 0xFD)
        add_rounded(slide, Inches(0.5), top, Inches(6), Inches(0.55), fill, ACCENT)
        add_text(slide, Inches(0.6), top + Inches(0.05), Inches(5.7), Inches(0.25),
                 layer, font_size=11, color=PRIMARY, bold=True)
        add_text(slide, Inches(0.6), top + Inches(0.28), Inches(5.7), Inches(0.25),
                 f"-> {desc}", font_size=10, color=DARK)
    add_text(slide, Inches(7), Inches(1.6), Inches(6), Inches(0.4), "TECHNOLOGY STACK", font_size=16, color=ACCENT, bold=True)
    tech_items = [
        "Language: Rust (memory-safe, zero-cost abstractions)",
        "UI Framework: Dioxus 0.7 (React-like, cross-platform)",
        "Networking: Axum (HTTP server) + rustls (TLS 1.3)",
        "Serialization: serde + JSON wire format",
        "Cryptography: rcgen (self-signed certs), TLS 1.3",
        "Storage: JSON files (trust, clipboard, notification history)",
        "Media: FFmpeg (camera), Piper (TTS), SenseVoice (STT)",
        "Build: Cargo + static linking for single-binary deployment"
    ]
    for i, item in enumerate(tech_items):
        add_text(slide, Inches(7.2), Inches(2.1 + i * 0.45), Inches(5.8), Inches(0.4),
                 f"[diams]  {item}", font_size=11, color=DARK)
    add_footer(slide, 5)
    return slide


def slide_workflow(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    SUCCESS = RGBColor(0x10, 0xB9, 0x81)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "System Workflow", "Data Flow & Interaction Patterns")
    add_text(slide, Inches(0.5), Inches(1.6), Inches(12.33), Inches(0.35),
             "End-to-End Data Flow", font_size=14, color=DARK, bold=True)
    workflow_steps = [
        ("1. Discovery", "Subnet scan (/24)\nDevice probes port 53327/53328\nAuto-detect peers on LAN", ACCENT),
        ("2. Pairing", "HTTPS link handshake\nTrust establishment\nPersistent trust store", SUCCESS),
        ("3. Clipboard", "1-second poll + SHA-256 hash\nPush to trusted peers via TLS\nApply on remote device", RGBColor(0x8B, 0x5C, 0xF6)),
        ("4. Notifications", "2-second poll for notifications\nMirror to trusted peers\nRemote reply support", RGBColor(0xF5, 0x9E, 0x0B)),
        ("5. File Transfer", "FastSwap protocol\nApproval popup (60s window)\n64KB chunked upload", RGBColor(0xEF, 0x44, 0x44)),
    ]
    for i, (title, desc, color) in enumerate(workflow_steps):
        left = Inches(0.5 + i * 2.5)
        top = Inches(2.2)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.8), top, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_text(slide, left + Inches(0.8), top + Inches(0.08), Inches(0.5), Inches(0.35),
                 str(i + 1), font_size=14, color=LIGHT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, left, top + Inches(0.6), Inches(2.3), Inches(0.3),
                 title, font_size=12, color=color, bold=True, align=PP_ALIGN.CENTER)
        lines = desc.split('\n')
        for j, line in enumerate(lines):
            add_text(slide, left, top + Inches(1.0 + j * 0.25), Inches(2.3), Inches(0.25),
                     line, font_size=9, color=SECONDARY, align=PP_ALIGN.CENTER)
        if i < len(workflow_steps) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left + Inches(2.1), top + Inches(0.1), Inches(0.4), Inches(0.3))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = SECONDARY
            arrow.line.fill.background()
    add_text(slide, Inches(0.5), Inches(4.2), Inches(12.33), Inches(0.35),
             "Network Topology", font_size=14, color=DARK, bold=True)
    topo_items = [
        ("Device A", "ClipboardManager (1s poll) -> SyncManager -> TLS Proxy (53328)"),
        ("Device B", "HTTP Server (53327) -> Apply clipboard -> Emit event to UI"),
        ("Discovery", "Bidirectional: HTTP (53327) + TLS (53328) scans"),
        ("FastSwap", "Sender: prepare -> confirm -> upload  |  Receiver: approval -> Downloads/"),
    ]
    for i, (label, flow) in enumerate(topo_items):
        row = i // 2
        col = i % 2
        left = Inches(0.5 + col * 6.15)
        top = Inches(4.7 + row * 1.1)
        add_rounded(slide, left, top, Inches(5.9), Inches(0.9),
                    RGBColor(0xFF, 0xFF, 0xFF), ACCENT)
        add_text(slide, left + Inches(0.2), top + Inches(0.08), Inches(5.4), Inches(0.3),
                 label, font_size=11, color=ACCENT, bold=True)
        add_text(slide, left + Inches(0.2), top + Inches(0.4), Inches(5.4), Inches(0.45),
                 flow, font_size=10, color=DARK)
    add_text(slide, Inches(0.5), Inches(7.0), Inches(12.33), Inches(0.3),
             "REST API: /api/ecosystem/v1/{info, clipboard/sync, notification/sync, notification/reply, pair/request, pair/untrust}",
             font_size=10, color=SECONDARY)
    add_footer(slide, 6)
    return slide


def slide_features(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    SUCCESS = RGBColor(0x10, 0xB9, 0x81)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Key Features", "What IGRIS Delivers")
    features = [
        ("Device Mesh", "Auto-discovery via subnet scan\nLink/trust handshake\n30s heartbeat, 120s timeout", ACCENT),
        ("Clipboard Sync", "Real-time text sync across devices\nSHA-256 loop-proof hashing\nTrust-gated delivery", SUCCESS),
        ("Notifications", "Mirror notifications across devices\nReply from any linked device\n100-entry history with read state", RGBColor(0x8B, 0x5C, 0xF6)),
        ("FastSwap Transfer", "FastSwap protocol\nTLS-encrypted, approval flow\n64KB chunks, folder preserved", RGBColor(0xF5, 0x9E, 0x0B)),
        ("Voice & Commands", "Optional voice assistant (STT/TTS)\nCommand processor + 27-tool registry\nPlugin system + setup manager", RGBColor(0xEF, 0x44, 0x44)),
        ("Cross-Platform", "Single native binary per OS\nWin32 / macOS / Linux\nPlatform abstraction traits", RGBColor(0x06, 0xB6, 0xD4)),
    ]
    for i, (title, desc, color) in enumerate(features):
        row = i // 3
        col = i % 3
        left = Inches(0.5 + col * 4.2)
        top = Inches(1.7 + row * 2.5)
        add_rounded(slide, left, top, Inches(3.9), Inches(2.2),
                    RGBColor(0xFF, 0xFF, 0xFF), color)
        add_rect(slide, left, top, Inches(3.9), Inches(0.5), color)
        add_text(slide, left + Inches(0.15), top + Inches(0.08), Inches(3.6), Inches(0.35),
                 title, font_size=14, color=LIGHT, bold=True)
        lines = desc.split('\n')
        for j, line in enumerate(lines):
            add_text(slide, left + Inches(0.15), top + Inches(0.6 + j * 0.35), Inches(3.6), Inches(0.3),
                     f"[bull] {line}", font_size=11, color=DARK)
    add_footer(slide, 7)
    return slide


def slide_conclusion(prs):
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    DARK = RGBColor(0x11, 0x18, 0x27)
    PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    SUCCESS = RGBColor(0x10, 0xB9, 0x81)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_title_bar(slide, "Conclusion", "Summary & Future Directions")
    add_text(slide, Inches(0.5), Inches(1.6), Inches(6), Inches(0.4), "SUMMARY", font_size=16, color=ACCENT, bold=True)
    summary_items = [
        "IGRIS delivers a fully local, serverless device mesh for Windows, macOS, and Linux",
        "All peer communication secured with TLS 1.3 and self-signed certificates",
        "Loop-proof clipboard sync, notification mirroring with reply, and FastSwap file transfer",
        "Optional voice assistant with offline STT/TTS and hybrid online mode",
        "Extensible plugin architecture supports voice interface and future features",
        "Single native binary deployment - no runtime dependencies or installation overhead"
    ]
    for i, item in enumerate(summary_items):
        add_text(slide, Inches(0.7), Inches(2.1 + i * 0.38), Inches(5.8), Inches(0.35),
                 f"[check]  {item}", font_size=11, color=DARK)
    add_text(slide, Inches(7), Inches(1.6), Inches(6), Inches(0.4), "FUTURE WORK", font_size=16, color=ACCENT, bold=True)
    future_items = [
        "Image/file clipboard sync beyond plain text",
        "Transfer history persistence across sessions",
        "Remote input and media control between devices",
        "Relay mode for non-LAN (cross-subnet / internet) peers",
        "Session handoff and shared AI memory between devices",
        "Mobile companion app (iOS / Android) integration"
    ]
    for i, item in enumerate(future_items):
        add_text(slide, Inches(7.2), Inches(2.1 + i * 0.38), Inches(5.8), Inches(0.35),
                 f"[arrow]  {item}", font_size=11, color=DARK)
    add_rounded(slide, Inches(0.5), Inches(5.8), Inches(12.33), Inches(0.9),
                RGBColor(0xDB, 0xEA, 0xFE), ACCENT)
    add_text(slide, Inches(0.7), Inches(5.9), Inches(12), Inches(0.35),
             "IGRIS redefines local device ecosystems: privacy-first, cross-platform, and completely free.",
             font_size=12, color=PRIMARY, bold=True)
    add_text(slide, Inches(0.7), Inches(6.25), Inches(12), Inches(0.35),
             "Licensed under MIT. Built with Rust for safety, speed, and portability.",
             font_size=11, color=SECONDARY)
    add_footer(slide, 8)
    return slide


def slide_references(prs):
    DARK = RGBColor(0x11, 0x18, 0x27)
    ACCENT = RGBColor(0x3B, 0x82, 0xF6)
    LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
    SECONDARY = RGBColor(0x64, 0x74, 0x8B)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, DARK)
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.2), RGBColor(0x0F, 0x17, 0x2A))
    add_rect(slide, Inches(0), Inches(1.2), prs.slide_width, Inches(0.06), ACCENT)
    add_text(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6),
             "References", font_size=32, color=LIGHT, bold=True)
    add_text(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.4),
             "Technologies, Protocols & Inspiration", font_size=14, color=SECONDARY)
    references = [
        "[1] Rust Programming Language - https://www.rust-lang.org",
        "[2] Dioxus Framework (v0.7) - https://dioxuslabs.com",
        "[3] Axum Web Framework - https://github.com/tokio-rs/axum",
        "[4] Rustls - Modern TLS library in Rust - https://github.com/rustls/rustls",
        "[5] LocalSend Protocol v2.0 - https://localsend.org",
        "[6] rcgen - Rust X.509 certificate generation - https://github.com/est31/rcgen",
        "[7] serde - Rust serialization framework - https://serde.rs",
        "[8] FFmpeg - Multimedia framework - https://ffmpeg.org",
        "[9] Piper TTS - Neural text-to-speech - https://github.com/rhasspy/piper",
        "[10] SenseVoice - Voice recognition model - Alibaba DAMO Academy",
        "[11] KDE Connect - Cross-device integration - https://kdeconnect.kde.org",
        "[12] Syncthing - Continuous file synchronization - https://syncthing.net",
        "[13] Open Neural Network Exchange (ONNX) - https://onnx.ai"
    ]
    for i, ref in enumerate(references):
        col = i // 8
        row = i % 8
        left = Inches(0.5 + col * 6.4)
        top = Inches(1.6 + row * 0.55)
        add_text(slide, left, top, Inches(6.2), Inches(0.45),
                 ref, font_size=11, color=RGBColor(0xCB, 0xD5, 0xE1))
    add_text(slide, Inches(0.5), Inches(6.8), Inches(12), Inches(0.3),
             "IGRIS: A Cross-Platform Device Ecosystem for the Modern LAN",
             font_size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    return slide
