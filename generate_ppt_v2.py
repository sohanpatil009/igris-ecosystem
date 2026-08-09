# generate_ppt_v2.py - IGRIS pitch deck: presented as a project we are building
import sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

DARK = RGBColor(0x11, 0x18, 0x27)
ACCENT = RGBColor(0x3B, 0x82, 0xF6)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
SECONDARY = RGBColor(0x64, 0x74, 0x8B)
PRIMARY = RGBColor(0x1A, 0x23, 0x7E)
SUCCESS = RGBColor(0x10, 0xB9, 0x81)

BULL = '\u25CF'
RARR = '\u2192'
CHECK = '\u2714'
ARROW = '\u27A4'
DIAMS = '\u25C6'


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_rounded(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, left, top, width, height, text, font_size=18, color=None, bold=False, align=PP_ALIGN.LEFT):
    if color is None:
        color = DARK
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Calibri"
    p.alignment = align
    return txBox


def add_title_bar(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.2), DARK)
    add_rect(slide, Inches(0), Inches(1.2), prs.slide_width, Inches(0.06), ACCENT)
    add_text(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6), title, font_size=28, color=LIGHT, bold=True)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.4), subtitle, font_size=14, color=SECONDARY)


def add_footer(slide, page_num, total=10):
    add_text(slide, Inches(0.5), Inches(7.1), Inches(12), Inches(0.3),
             f"IGRIS - Cross-Platform Device Ecosystem    |    Page {page_num}/{total}",
             font_size=10, color=SECONDARY)


# ---------- Slide 1: Title ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_rect(slide, Inches(0), Inches(2.5), prs.slide_width, Inches(0.08), ACCENT)
add_text(slide, Inches(0.5), Inches(1.0), Inches(12), Inches(0.8), "IGRIS", font_size=60, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.5), Inches(1.8), Inches(12), Inches(0.6), "One Device Ecosystem for Desktop & Mobile", font_size=32, color=LIGHT, align=PP_ALIGN.CENTER)
add_rect(slide, Inches(4.5), Inches(3.5), Inches(4.33), Inches(0.05), ACCENT)
items = ["Universal Clipboard Sync", "Cross-Device Notification Mirror",
         "FastSwap File Transfer (P2P, TLS-encrypted)", "Device Mesh for Windows, macOS, Linux, Android, iOS"]
for i, item in enumerate(items):
    add_text(slide, Inches(3.5), Inches(3.8 + i * 0.45), Inches(6.33), Inches(0.4),
             f"{BULL}  {item}", font_size=16, color=RGBColor(0x94, 0xA3, 0xB8), align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.5), Inches(6.3), Inches(12), Inches(0.4),
         "We are building a Rust-powered, serverless, account-free ecosystem for your devices",
         font_size=14, color=SECONDARY, align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.5), Inches(6.75), Inches(12), Inches(0.35),
         "No cloud. No accounts. Just your devices, on your network.",
         font_size=12, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

# ---------- Slide 2: Vision ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Vision", "What We Are Building")
facts = [
    "One app across every device you own - clipboard, notifications, and files follow you",
    "A private device mesh on your LAN - serverless, account-free, no internet required",
    "Zero-config setup - devices on the same network find each other automatically",
    "Trust-gated by design - data flows only between linked devices, encrypted end-to-end",
    "FastSwap file transfer - send files between any of your devices with one approval",
    "An optional voice assistant that understands commands and works without the cloud",
]
for i, fact in enumerate(facts):
    row = i // 2
    col = i % 2
    left = Inches(0.5 + col * 6.15)
    top = Inches(1.7 + row * 1.1)
    add_rounded(slide, left, top, Inches(5.9), Inches(0.95), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE2, 0xE8, 0xF0))
    add_text(slide, left + Inches(0.2), top + Inches(0.1), Inches(5.4), Inches(0.8),
             f"{CHECK}  {fact}", font_size=12, color=DARK)
add_text(slide, Inches(0.5), Inches(5.4), Inches(12.33), Inches(0.5),
         "IGRIS will be the single app that unifies every screen you own.",
         font_size=16, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)
add_footer(slide, 2)

# ---------- Slide 3: Problem ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "The Problem", "Why We Are Building IGRIS")
add_text(slide, Inches(0.5), Inches(1.5), Inches(6), Inches(0.35), "TODAY'S DEVICE WORKFLOWS ARE BROKEN", font_size=14, color=RGBColor(0xEF, 0x44, 0x44), bold=True)
problems = [
    f"{DIAMS}  Files, clipboard, and notifications are trapped on one device",
    f"{DIAMS}  Moving a file between phone and laptop means cables, cloud uploads, or new accounts",
    f"{DIAMS}  Cloud sync copies your data to someone else's servers",
    f"{DIAMS}  Vendor sync tools are walled gardens - AirDrop, Quick Share, and KDE Connect don't talk to each other",
    f"{DIAMS}  Nothing ties a Windows desktop and an Android phone together without signing up",
]
for i, line in enumerate(problems):
    add_text(slide, Inches(0.7), Inches(2.0 + i * 0.6), Inches(5.9), Inches(0.55),
             line, font_size=12, color=DARK)
add_text(slide, Inches(6.8), Inches(1.5), Inches(6), Inches(0.35), "OUR ANSWER", font_size=14, color=ACCENT, bold=True)
answers = [
    f"{CHECK}  One open ecosystem that works across every major OS",
    f"{CHECK}  Peer-to-peer over your own network - data never leaves it",
    f"{CHECK}  Works with no accounts and no internet connection",
    f"{CHECK}  Native experience on every platform, one shared core",
    f"{CHECK}  Private by default: nothing is collected, uploaded, or sold",
]
for i, line in enumerate(answers):
    add_text(slide, Inches(7.0), Inches(2.0 + i * 0.6), Inches(5.9), Inches(0.55),
             line, font_size=12, color=DARK)
add_footer(slide, 3)

# ---------- Slide 4: Core Features ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Core Features", "What IGRIS Will Do")
features = [
    ("Device Mesh", f"{BULL} Automatic discovery of your devices\n{BULL} Link & trust handshake\n{BULL} Live presence across the network\n{BULL} Works with zero configuration", ACCENT),
    ("Universal Clipboard", f"{BULL} Copy on one device, paste on another\n{BULL} Loop-proof, hash-diffed sync\n{BULL} Trust-gated delivery\n{BULL} Plain text first, more formats next", SUCCESS),
    ("Notification Mirror", f"{BULL} See any device's notifications\n{BULL} Reply from whichever device is near\n{BULL} Read history across devices\n{BULL} Platform-native look & feel", RGBColor(0x8B, 0x5C, 0xF6)),
    ("FastSwap", f"{BULL} Fast, encrypted file transfer\n{BULL} Approval popup on the receiver\n{BULL} Folders preserved, large files\n{BULL} Cancel & resume", RGBColor(0xF5, 0x9E, 0x0B)),
    ("Cross-Platform", f"{BULL} Desktop: Windows, macOS, Linux\n{BULL} Mobile: Android & iOS\n{BULL} One Rust core on every platform\n{BULL} Native UI per OS", RGBColor(0x06, 0xB6, 0xD4)),
    ("Voice Assistant", f"{BULL} Offline voice commands\n{BULL} On-device speech recognition\n{BULL} Works without the cloud\n{BULL} Extensible command system", RGBColor(0xEF, 0x44, 0x44)),
]
for i, (title, desc, color) in enumerate(features):
    row = i // 3
    col = i % 3
    left = Inches(0.5 + col * 4.2)
    top = Inches(1.7 + row * 2.5)
    add_rounded(slide, left, top, Inches(3.9), Inches(2.2), RGBColor(0xFF, 0xFF, 0xFF), color)
    add_rect(slide, left, top, Inches(3.9), Inches(0.5), color)
    add_text(slide, left + Inches(0.15), top + Inches(0.08), Inches(3.6), Inches(0.35),
             title, font_size=14, color=LIGHT, bold=True)
    lines = desc.split('\n')
    for j, line in enumerate(lines):
        add_text(slide, left + Inches(0.15), top + Inches(0.55 + j * 0.4), Inches(3.6), Inches(0.35),
                 line, font_size=10, color=DARK)
add_footer(slide, 4)

# ---------- Slide 5: Architecture ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Architecture", "One Core, Many Surfaces")
add_text(slide, Inches(0.5), Inches(1.5), Inches(12.33), Inches(0.35),
         "A shared Rust core holds all the intelligence; each platform gets a thin native surface on top.",
         font_size=13, color=PRIMARY, bold=True)
arch = [
    ("Shared Rust Core", "Device mesh, discovery, sync engine, TLS encryption, file transfer, crypto - written once in Rust, compiled natively for every platform"),
    ("Desktop UI", "Windows, macOS, and Linux apps with a native, responsive interface"),
    ("Mobile UI", "Android app in Kotlin, iOS app in Swift - thin UI layers over the same core"),
    ("UniFFI Bindings", "Safe, generated FFI bindings that expose the Rust core to Kotlin and Swift"),
    ("Platform Abstraction", "Clipboard, notifications, and system services implemented per OS"),
    ("Plugin System", "Modular core that grows: voice, commands, media, remote control"),
]
for i, (name, desc) in enumerate(arch):
    row = i // 2
    col = i % 2
    left = Inches(0.5 + col * 6.15)
    top = Inches(2.0 + row * 1.25)
    add_rounded(slide, left, top, Inches(5.9), Inches(1.05), RGBColor(0xDB, 0xEA, 0xFE), ACCENT)
    add_text(slide, left + Inches(0.2), top + Inches(0.08), Inches(5.4), Inches(0.3),
             name, font_size=12, color=PRIMARY, bold=True)
    add_text(slide, left + Inches(0.2), top + Inches(0.4), Inches(5.4), Inches(0.6),
             desc, font_size=10, color=DARK)
add_text(slide, Inches(0.5), Inches(6.0), Inches(12.33), Inches(0.4),
         f"{ARROW}  Write the ecosystem once in Rust - reuse it on every desktop and mobile platform.",
         font_size=12, color=ACCENT, bold=True)
add_footer(slide, 5)

# ---------- Slide 6: Mobile Strategy ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Mobile Strategy", "Native Kotlin + Swift on a Shared Rust Core")
add_text(slide, Inches(0.5), Inches(1.5), Inches(12.33), Inches(0.35),
         "All device logic lives once in the Rust core; each OS only carries its own thin UI.",
         font_size=13, color=PRIMARY, bold=True)
mobile = [
    ("Rust Core (shared)", f"{BULL} Mesh, discovery, sync, TLS, transfer\n{BULL} Runs on Android & iOS via UniFFI\n{BULL} Same behavior on every device", ACCENT),
    ("Android App", f"{BULL} Native UI written in Kotlin\n{BULL} Calls the core through UniFFI bindings\n{BULL} Feels like a native Android app", SUCCESS),
    ("iOS App", f"{BULL} Native UI written in Swift\n{BULL} Same Rust core, same bindings\n{BULL} Feels like a native iOS app", RGBColor(0x8B, 0x5C, 0xF6)),
]
for i, (title, desc, color) in enumerate(mobile):
    left = Inches(0.5 + i * 4.2)
    top = Inches(2.0)
    add_rounded(slide, left, top, Inches(3.9), Inches(2.0), RGBColor(0xFF, 0xFF, 0xFF), color)
    add_rect(slide, left, top, Inches(3.9), Inches(0.5), color)
    add_text(slide, left + Inches(0.15), top + Inches(0.08), Inches(3.6), Inches(0.35),
             title, font_size=14, color=LIGHT, bold=True)
    lines = desc.split('\n')
    for j, line in enumerate(lines):
        add_text(slide, left + Inches(0.15), top + Inches(0.55 + j * 0.4), Inches(3.6), Inches(0.35),
                 line, font_size=10, color=DARK)
add_rounded(slide, Inches(0.5), Inches(4.3), Inches(12.33), Inches(2.1), RGBColor(0xFF, 0xF7, 0xE6), RGBColor(0xF5, 0x9E, 0x0B))
add_text(slide, Inches(0.7), Inches(4.45), Inches(11.9), Inches(0.35), "WHY NATIVE KOTLIN + SWIFT INSTEAD OF FLUTTER?", font_size=14, color=RGBColor(0xB4, 0x69, 0x06), bold=True)
why = [
    f"{BULL}  OS-specific code has to be written in the platform language anyway - the Rust core is already cross-platform",
    f"{BULL}  Flutter would add a third UI layer without removing any native OS work",
    f"{BULL}  Native apps are smaller, faster to start, and feel at home on each OS",
    f"{BULL}  One shared Rust core + thin native UI is the simplest way to cover all five platforms",
]
for i, line in enumerate(why):
    add_text(slide, Inches(0.9), Inches(4.85 + i * 0.4), Inches(11.7), Inches(0.35),
             line, font_size=11, color=DARK)
add_text(slide, Inches(0.5), Inches(6.5), Inches(12.33), Inches(0.4),
         f"{ARROW}  Result: clipboard, notifications, and files flow between your phone and your desktop - in both directions.",
         font_size=12, color=ACCENT, bold=True)
add_footer(slide, 6)

# ---------- Slide 7: Networking & Security ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Networking & Security", "Private by Design")
add_text(slide, Inches(0.5), Inches(1.5), Inches(6), Inches(0.35), "HOW IT WORKS", font_size=14, color=ACCENT, bold=True)
net = [
    f"{DIAMS}  LAN-first: all peer traffic stays on your local network",
    f"{DIAMS}  Zero-config discovery: devices on the subnet find each other automatically",
    f"{DIAMS}  Link handshake establishes trust between two devices",
    f"{DIAMS}  Data flows only between linked devices",
    f"{DIAMS}  Optional relay later, for cross-network use",
]
for i, line in enumerate(net):
    add_text(slide, Inches(0.7), Inches(2.0 + i * 0.6), Inches(5.9), Inches(0.5),
             line, font_size=12, color=DARK)
add_text(slide, Inches(6.8), Inches(1.5), Inches(6), Inches(0.35), "SECURITY", font_size=14, color=ACCENT, bold=True)
sec = [
    f"{CHECK}  All peer traffic encrypted with TLS 1.3",
    f"{CHECK}  Self-signed certificates generated per device",
    f"{CHECK}  Accept-on-LAN trust model - you approve every link",
    f"{CHECK}  No cloud, no accounts, no telemetry",
    f"{CHECK}  Your data never leaves your network",
]
for i, line in enumerate(sec):
    add_text(slide, Inches(7.0), Inches(2.0 + i * 0.6), Inches(5.9), Inches(0.5),
             line, font_size=12, color=DARK)
add_rounded(slide, Inches(0.5), Inches(5.4), Inches(12.33), Inches(1.0), RGBColor(0xFF, 0xF7, 0xE6), RGBColor(0xF5, 0x9E, 0x0B))
add_text(slide, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.7),
         f"{BULL}  Privacy is the product: IGRIS never sees your data, because your data never leaves your house.",
         font_size=14, color=RGBColor(0xB4, 0x69, 0x06), bold=True)
add_footer(slide, 7)

# ---------- Slide 8: Tech Stack ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Tech Stack", "The Building Blocks")
add_text(slide, Inches(0.5), Inches(1.5), Inches(6), Inches(0.35), "PLANNED STACK", font_size=14, color=ACCENT, bold=True)
tech = [
    "Core: Rust (memory-safe, fast, portable)",
    "Async runtime: tokio",
    "HTTP servers: axum",
    "TLS: rustls + self-signed certs (rcgen)",
    "Desktop UI: Dioxus",
    "Mobile UI: Kotlin (Android) + Swift (iOS)",
    "Core-to-mobile bridge: UniFFI bindings",
    "Voice: on-device STT/TTS, offline-first",
]
for i, line in enumerate(tech):
    add_text(slide, Inches(0.7), Inches(1.95 + i * 0.38), Inches(5.9), Inches(0.35),
             f"{DIAMS}  {line}", font_size=11, color=DARK)
add_text(slide, Inches(6.8), Inches(1.5), Inches(6), Inches(0.35), "DESIGN PRINCIPLES", font_size=14, color=ACCENT, bold=True)
principles = [
    f"{CHECK}  Write once in Rust, run everywhere",
    f"{CHECK}  Offline-first: every feature works without the internet",
    f"{CHECK}  No accounts, no cloud, no telemetry",
    f"{CHECK}  Open-source and MIT licensed",
    f"{CHECK}  Single binary per platform",
    f"{CHECK}  Modular core, plugin-friendly",
]
for i, line in enumerate(principles):
    add_text(slide, Inches(7.0), Inches(1.95 + i * 0.38), Inches(5.9), Inches(0.35),
             f"{DIAMS}  {line}", font_size=11, color=DARK)
add_footer(slide, 8)

# ---------- Slide 9: Roadmap ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, LIGHT)
add_title_bar(slide, "Roadmap", "How We Get There")
phases = [
    ("Phase 1 - Desktop Foundation", "Windows, macOS, Linux apps", "Device mesh, discovery, clipboard sync, notification mirror, FastSwap file transfer, voice assistant", ACCENT),
    ("Phase 2 - Mobile Apps", "Android (Kotlin) + iOS (Swift)", "Thin native UIs over the shared Rust core via UniFFI; clipboard, notifications, and files flow between phone and desktop", SUCCESS),
    ("Phase 3 - Beyond the LAN", "Optional relay mode", "Cross-network sync, richer clipboard formats (images, files), remote device control, shared AI memory", RGBColor(0x8B, 0x5C, 0xF6)),
]
for i, (name, when, what, color) in enumerate(phases):
    top = Inches(1.6 + i * 1.65)
    add_rounded(slide, Inches(0.5), top, Inches(12.33), Inches(1.45), RGBColor(0xFF, 0xFF, 0xFF), color)
    add_rect(slide, Inches(0.5), top, Inches(2.6), Inches(1.45), color)
    add_text(slide, Inches(0.65), top + Inches(0.45), Inches(2.3), Inches(0.6),
             name, font_size=13, color=LIGHT, bold=True)
    add_text(slide, Inches(3.4), top + Inches(0.12), Inches(9.2), Inches(0.35),
             when, font_size=12, color=PRIMARY, bold=True)
    add_text(slide, Inches(3.4), top + Inches(0.5), Inches(9.2), Inches(0.8),
             what, font_size=11, color=DARK)
add_text(slide, Inches(0.5), Inches(6.6), Inches(12.33), Inches(0.4),
         f"{ARROW}  Mobile is the key next milestone: the same Rust core, now on the devices people carry.",
         font_size=13, color=ACCENT, bold=True)
add_footer(slide, 9)

# ---------- Slide 10: Summary ----------
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_rect(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.2), RGBColor(0x0F, 0x17, 0x2A))
add_rect(slide, Inches(0), Inches(1.2), prs.slide_width, Inches(0.06), ACCENT)
add_text(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.6), "Summary", font_size=32, color=LIGHT, bold=True)
summary = [
    f"{CHECK}  We are building one private device ecosystem - Windows, macOS, Linux, Android, iOS",
    f"{CHECK}  One Rust core, shared everywhere; native UI on every platform",
    f"{CHECK}  Clipboard, notifications, and files follow you between devices",
    f"{CHECK}  Encrypted peer-to-peer on your LAN - no cloud, no accounts",
    f"{CHECK}  Mobile apps in Kotlin and Swift, powered by the same core via UniFFI",
    f"{CHECK}  Open-source, MIT licensed, private by design",
    f"{CHECK}  Next milestone: native mobile apps that bring the desktop mesh into your pocket",
]
for i, line in enumerate(summary):
    add_text(slide, Inches(0.7), Inches(1.7 + i * 0.5), Inches(12), Inches(0.4),
             line, font_size=13, color=RGBColor(0xCB, 0xD5, 0xE1))
add_rect(slide, Inches(0), Inches(6.3), prs.slide_width, Inches(0.06), ACCENT)
add_text(slide, Inches(0.5), Inches(6.5), Inches(12), Inches(0.4),
         "Your devices, finally working as one - and nothing leaves your network.",
         font_size=14, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.5), Inches(6.95), Inches(12), Inches(0.3),
         "Built on Rust for safety, speed, and portability.",
         font_size=11, color=SECONDARY, align=PP_ALIGN.CENTER)

output_path = r"D:\ecosystem\igrisv4\IGRIS_Presentation_v2.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
